from pathlib import Path

import openpyxl
import pytest

from mcnp_report.ai import local_result
from mcnp_report.parser import parse_output
from mcnp_report.report import write_report


def test_workbook_content_and_charts(tmp_path):
    source = Path(r"D:\mcnpproject\1.out")
    if not source.is_file(): pytest.skip("external sample unavailable")
    result = parse_output(source)
    destination = tmp_path / "report.xlsx"
    try:
        write_report(result, local_result(result), destination)
    finally: result.close()
    workbook = openpyxl.load_workbook(destination, read_only=False, data_only=False)
    assert workbook.sheetnames == ["总览", "Tally结果", "统计诊断", "粒子与单元", "核数据库", "输入清单", "解析审计"]
    assert workbook["总览"]["B10"].value == pytest.approx(4.95542e-4)
    assert workbook["总览"]["B11"].value == pytest.approx(0.0035)
    assert workbook["Tally结果"].max_row == 2007
    assert isinstance(workbook["Tally结果"]["H5"].value, (int, float))
    assert workbook["Tally结果"]["K2007"].value == "总计"
    assert len(workbook["总览"]._charts) == 2
    workbook.close()


def test_tally_sheet_split(monkeypatch, tmp_path):
    source = Path(r"D:\mcnpproject\1.out")
    if not source.is_file(): pytest.skip("external sample unavailable")
    import mcnp_report.report as report_module
    monkeypatch.setattr(report_module, "TALLY_ROWS_PER_SHEET", 1000)
    result = parse_output(source)
    destination = tmp_path / "split.xlsx"
    try: report_module.write_report(result, local_result(result), destination)
    finally: result.close()
    workbook = openpyxl.load_workbook(destination, read_only=True)
    assert workbook.sheetnames[:4] == ["总览", "Tally结果_01", "Tally结果_02", "Tally结果_03"]
    workbook.close()
